import os
import zipfile
from pathlib import Path
from typing import Optional, List, Dict, Any
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.pdfgen import canvas
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from app.config import EXPORT_DIR
from app.models import QuestionPaper, PaperSet, Question
from app.services.template_service import get_template


class NumberedCanvas(canvas.Canvas):
    """Adds 'Page X of Y' numbering to ReportLab PDFs."""
    def __init__(self, *args, **kwargs):
        super(NumberedCanvas, self).__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            canvas.Canvas.showPage(self)
        canvas.Canvas.save(self)

    def draw_page_number(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 9)
        self.setFillColor(colors.HexColor("#4b5563"))
        text = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(A4[0] - 40, 30, text)
        self.drawString(40, 30, "SmartQP – Confidential Examination Paper")
        self.setStrokeColor(colors.HexColor("#d1d5db"))
        self.setLineWidth(0.5)
        self.line(40, 42, A4[0] - 40, 42)
        self.restoreState()


def export_set_to_pdf(
    paper: QuestionPaper,
    paper_set: PaperSet,
    include_tags: bool = False,
    template_id: str = "autonomous_obe",
    custom_template: Optional[Dict[str, Any]] = None
) -> str:
    """Generates a polished academic PDF examination paper based on the selected Institute Template."""
    tpl = get_template(template_id, custom_template)
    
    filename = f"Paper_{paper.id}_{paper_set.set_label.replace(' ', '_')}_{template_id}.pdf"
    file_path = str(EXPORT_DIR / filename)

    doc = SimpleDocTemplate(
        file_path,
        pagesize=A4,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=50
    )

    styles = getSampleStyleSheet()
    
    inst_style = ParagraphStyle(
        'InstHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=14,
        leading=18,
        alignment=1, # Center
        textColor=colors.HexColor("#0f172a")
    )
    
    sub_header_style = ParagraphStyle(
        'SubHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10,
        leading=14,
        alignment=1,
        textColor=colors.HexColor("#334155")
    )

    exam_title_style = ParagraphStyle(
        'ExamTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=15,
        alignment=1,
        textColor=colors.HexColor("#1e293b")
    )

    set_badge_style = ParagraphStyle(
        'SetBadge',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=13,
        alignment=1,
        textColor=colors.HexColor("#4338ca")
    )

    meta_style = ParagraphStyle(
        'MetaStyle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#1e293b")
    )

    meta_bold = ParagraphStyle(
        'MetaBold',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#1e293b")
    )

    instruc_style = ParagraphStyle(
        'Instruction',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=8.5,
        leading=11.5,
        textColor=colors.HexColor("#475569")
    )

    q_text_style = ParagraphStyle(
        'QText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor("#0f172a")
    )

    col_center_style = ParagraphStyle(
        'ColCenter',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=12,
        alignment=1, # Center
        textColor=colors.HexColor("#0f172a")
    )

    story = []

    # 1. Optional USN / Register Number Grid (if template enables it)
    if tpl.get("show_usn_box", True):
        usn_boxes = [Paragraph("<b>USN / Reg. No.</b>", meta_bold)]
        # Add 10 individual character boxes for student register numbers
        for _ in range(10):
            usn_boxes.append(Paragraph("", meta_style))
        usn_table = Table([usn_boxes], colWidths=[95] + [16]*10)
        usn_table.setStyle(TableStyle([
            ('BOX', (0,0), (-1,-1), 0.75, colors.HexColor("#475569")),
            ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#94a3b8")),
            ('BACKGROUND', (0,0), (0,0), colors.HexColor("#f1f5f9")),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('TOPPADDING', (0,0), (-1,-1), 3),
            ('BOTTOMPADDING', (0,0), (-1,-1), 3),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ]))
        story.append(usn_table)
        story.append(Spacer(1, 6))

    # 2. Header Information (From Template or Subject)
    inst_name = tpl.get("header_title") or (paper.subject.institute_name if paper.subject else "INSTITUTION OF HIGHER EDUCATION")
    dept_name = paper.subject.department_name if paper.subject else "DEPARTMENT OF COMPUTER SCIENCE & ENGINEERING"
    exam_name = paper.examination_name or tpl.get("header_subtitle") or "SEMESTER END EXAMINATION"

    story.append(Paragraph(inst_name.upper(), inst_style))
    story.append(Spacer(1, 2))
    story.append(Paragraph(dept_name.upper(), sub_header_style))
    story.append(Spacer(1, 3))
    story.append(Paragraph(exam_name.upper(), exam_title_style))
    story.append(Spacer(1, 5))

    # Set Badge
    set_box_data = [[Paragraph(f"• {paper_set.set_label.upper()} •", set_badge_style)]]
    set_box_table = Table(set_box_data, colWidths=[120])
    set_box_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#eef2ff")),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor("#6366f1")),
        ('TOPPADDING', (0,0), (-1,-1), 3),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
    ]))
    story.append(set_box_table)
    story.append(Spacer(1, 6))

    # 3. Meta Table (Course, Subject, Max Marks, Time)
    subj_code_str = f" ({paper.subject.subject_code})" if paper.subject and paper.subject.subject_code else ""
    meta_table_data = [
        [
            Paragraph(f"<b>Course / Program:</b> {paper.subject.course_name if paper.subject else 'B.Tech'}", meta_style),
            Paragraph(f"<b>Duration:</b> {paper.duration_minutes or 180} Minutes", meta_style),
        ],
        [
            Paragraph(f"<b>Subject:</b> {paper.subject.subject_name if paper.subject else 'Subject'}{subj_code_str}", meta_style),
            Paragraph(f"<b>Max. Marks:</b> {paper_set.total_marks}", meta_bold),
        ],
        [
            Paragraph(f"<b>Semester / Year:</b> {paper.subject.semester or ''} {paper.subject.academic_year or ''}", meta_style),
            Paragraph(f"<b>Total Questions:</b> {len(paper_set.questions)}", meta_style),
        ]
    ]

    meta_table = Table(meta_table_data, colWidths=[335, 185])
    meta_table.setStyle(TableStyle([
        ('BOX', (0,0), (-1,-1), 0.75, colors.HexColor("#94a3b8")),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#e2e8f0")),
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#f8fafc")),
        ('TOPPADDING', (0,0), (-1,-1), 3),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3),
        ('LEFTPADDING', (0,0), (-1,-1), 6),
        ('RIGHTPADDING', (0,0), (-1,-1), 6),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 6))

    # 4. Instructions
    instructions_text = tpl.get("instructions") or paper.instructions or "1. Answer all questions. 2. Figures to the right indicate full marks."
    story.append(Paragraph(f"<b>Instructions:</b> {instructions_text.replace(chr(10), '<br/>')}", instruc_style))
    story.append(Spacer(1, 8))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#334155"), spaceAfter=8))

    # 5. Questions Table Header
    show_obe = tpl.get("show_obe_columns", False) or include_tags

    if show_obe:
        # NBA/NAAC Outcome Based Education (OBE) layout with CO and Bloom's Level
        q_table_data = [
            [
                Paragraph("<b>Q.No</b>", col_center_style),
                Paragraph("<b>Question Description</b>", meta_bold),
                Paragraph("<b>Marks</b>", col_center_style),
                Paragraph("<b>CO</b>", col_center_style),
                Paragraph("<b>BTL</b>", col_center_style),
            ]
        ]
        col_widths = [35, 365, 42, 38, 40]
    else:
        # Standard layout
        q_table_data = [
            [
                Paragraph("<b>Q.No</b>", col_center_style),
                Paragraph("<b>Question Description</b>", meta_bold),
                Paragraph("<b>Marks</b>", ParagraphStyle('HRight', parent=meta_bold, alignment=2))
            ]
        ]
        col_widths = [40, 430, 50]

    questions = sorted(paper_set.questions, key=lambda q: q.question_number)
    for q in questions:
        q_content = Paragraph(q.question_text, q_text_style)

        if show_obe:
            q_table_data.append([
                Paragraph(f"<b>Q{q.question_number}</b>", col_center_style),
                q_content,
                Paragraph(f"[{q.marks}]", col_center_style),
                Paragraph(f"{q.co_level or 'CO1'}", col_center_style),
                Paragraph(f"{q.bloom_level[:4] if q.bloom_level else 'L2'}", col_center_style),
            ])
        else:
            q_table_data.append([
                Paragraph(f"<b>Q{q.question_number}.</b>", meta_style),
                q_content,
                Paragraph(f"[{q.marks}]", ParagraphStyle('MRight', parent=meta_bold, alignment=2))
            ])

    q_table = Table(q_table_data, colWidths=col_widths)
    q_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#f1f5f9")),
        ('BOX', (0,0), (-1,-1), 0.75, colors.HexColor("#0f172a")),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('TOPPADDING', (0,0), (-1,-1), 5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 5),
        ('LEFTPADDING', (0,0), (-1,-1), 4),
        ('RIGHTPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(q_table)

    # 6. Optional Sign-Off Blocks from Template
    sign_blocks = tpl.get("sign_off_blocks", [])
    if sign_blocks:
        story.append(Spacer(1, 24))
        sign_cols = [Paragraph(f"___________________<br/><b>{block}</b>", ParagraphStyle('Sign', parent=meta_style, alignment=1)) for block in sign_blocks]
        sign_table = Table([sign_cols], colWidths=[520 // len(sign_blocks)] * len(sign_blocks))
        sign_table.setStyle(TableStyle([
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ]))
        story.append(sign_table)

    story.append(Spacer(1, 16))
    story.append(Paragraph("--- END OF QUESTION PAPER ---", ParagraphStyle('EndMsg', parent=styles['Normal'], alignment=1, fontSize=8.5, textColor=colors.HexColor("#64748b"))))

    doc.build(story, canvasmaker=NumberedCanvas)
    return file_path


def export_set_to_docx(
    paper: QuestionPaper,
    paper_set: PaperSet,
    include_tags: bool = False,
    template_id: str = "autonomous_obe",
    custom_template: Optional[Dict[str, Any]] = None
) -> str:
    """Generates a clean formatted Word (.docx) document strictly following the chosen Institute Template."""
    tpl = get_template(template_id, custom_template)
    filename = f"Paper_{paper.id}_{paper_set.set_label.replace(' ', '_')}_{template_id}.docx"
    file_path = str(EXPORT_DIR / filename)

    doc = docx.Document()

    # 1. Optional USN / Register Number Grid
    if tpl.get("show_usn_box", True):
        usn_table = doc.add_table(rows=1, cols=11)
        usn_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        usn_table.cell(0, 0).text = "USN / Reg No:"
        usn_table.cell(0, 0).paragraphs[0].runs[0].bold = True
        for col_i in range(1, 11):
            usn_table.cell(0, col_i).text = " "
        doc.add_paragraph()

    # 2. Document Header
    inst_name = tpl.get("header_title") or (paper.subject.institute_name if paper.subject else "INSTITUTION OF HIGHER EDUCATION")
    inst_p = doc.add_paragraph()
    inst_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    inst_run = inst_p.add_run(inst_name.upper())
    inst_run.bold = True
    inst_run.font.size = Pt(13)

    dept_p = doc.add_paragraph()
    dept_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    dept_run = dept_p.add_run(paper.subject.department_name.upper() if paper.subject else "DEPARTMENT")
    dept_run.bold = True
    dept_run.font.size = Pt(10.5)

    exam_title = paper.examination_name or tpl.get("header_subtitle") or "SEMESTER EXAMINATION"
    exam_p = doc.add_paragraph()
    exam_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    exam_run = exam_p.add_run(f"{exam_title.upper()} - [{paper_set.set_label.upper()}]")
    exam_run.bold = True
    exam_run.font.size = Pt(11)
    exam_run.font.color.rgb = RGBColor(67, 56, 202)

    # 3. Details table
    table = doc.add_table(rows=3, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    table.cell(0, 0).text = f"Course: {paper.subject.course_name if paper.subject else ''}"
    table.cell(0, 1).text = f"Duration: {paper.duration_minutes or 180} Minutes"
    table.cell(1, 0).text = f"Subject: {paper.subject.subject_name if paper.subject else ''} ({paper.subject.subject_code or ''})"
    table.cell(1, 1).text = f"Maximum Marks: {paper_set.total_marks}"
    table.cell(2, 0).text = f"Semester/Year: {paper.subject.semester or ''} {paper.subject.academic_year or ''}"
    table.cell(2, 1).text = f"Total Questions: {len(paper_set.questions)}"

    doc.add_paragraph()
    
    # 4. Instructions
    instruc_p = doc.add_paragraph()
    instructions_text = tpl.get("instructions") or paper.instructions or "Answer all questions. Figures to the right indicate full marks."
    instruc_p.add_run(f"Instructions: {instructions_text}").italic = True

    doc.add_paragraph()

    # 5. Questions Table
    show_obe = tpl.get("show_obe_columns", False) or include_tags
    num_cols = 5 if show_obe else 3

    q_table = doc.add_table(rows=1, cols=num_cols)
    q_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = q_table.rows[0].cells
    
    hdr_cells[0].text = "Q.No"
    hdr_cells[1].text = "Question Description"
    hdr_cells[2].text = "Marks"
    if show_obe:
        hdr_cells[3].text = "CO"
        hdr_cells[4].text = "Bloom's"

    for cell in hdr_cells:
        if cell.paragraphs and cell.paragraphs[0].runs:
            cell.paragraphs[0].runs[0].bold = True

    questions = sorted(paper_set.questions, key=lambda q: q.question_number)
    for q in questions:
        row_cells = q_table.add_row().cells
        row_cells[0].text = f"Q{q.question_number}."
        row_cells[1].text = q.question_text
        row_cells[2].text = f"[{q.marks}]"
        if show_obe:
            row_cells[3].text = str(q.co_level or "CO1")
            row_cells[4].text = str(q.bloom_level or "Understand")

    # 6. Optional Sign-Off Blocks
    sign_blocks = tpl.get("sign_off_blocks", [])
    if sign_blocks:
        doc.add_paragraph()
        doc.add_paragraph()
        sign_table = doc.add_table(rows=1, cols=len(sign_blocks))
        sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for idx, block in enumerate(sign_blocks):
            cell = sign_table.cell(0, idx)
            cell.text = f"____________________\n{block}"
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()
    end_p = doc.add_paragraph("--- END OF QUESTION PAPER ---")
    end_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(file_path)
    return file_path


def export_all_sets_bundle(
    paper: QuestionPaper,
    fmt: str = "pdf",
    include_tags: bool = False,
    template_id: str = "autonomous_obe",
    custom_template: Optional[Dict[str, Any]] = None
) -> str:
    """Exports all sets as a single ZIP archive formatted according to the selected Institute Template."""
    zip_filename = f"Paper_{paper.id}_All_Sets_{template_id}_{fmt.upper()}.zip"
    zip_path = str(EXPORT_DIR / zip_filename)

    generated_files = []
    for s in paper.sets:
        if fmt.lower() == "pdf":
            fpath = export_set_to_pdf(paper, s, include_tags=include_tags, template_id=template_id, custom_template=custom_template)
        else:
            fpath = export_set_to_docx(paper, s, include_tags=include_tags, template_id=template_id, custom_template=custom_template)
        generated_files.append(fpath)

    with zipfile.ZipFile(zip_path, 'w') as zipf:
        for f in generated_files:
            zipf.write(f, arcname=os.path.basename(f))

    return zip_path
